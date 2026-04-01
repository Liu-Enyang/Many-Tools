from __future__ import annotations

from pathlib import Path
import re
import json
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook



CHECKLIST_SHEET_NAME = "UT仕様書チェックシート"
CHECKLIST_START_ROW = 26
CHECKLIST_END_ROW = 116


CHECKLIST_MATCH_RULES: List[Dict[str, Any]] = [
    {
        "nos": ["2-1", "2-2", "2-6"],
        "keywords": ["画面レイアウト", "画面構成", "ボタン配置", "崩れない"],
    },
    {
        "nos": ["2-7", "2-9"],
        "keywords": ["初期表示", "想定データ", "表示されること", "初期値"],
    },
    {
        "nos": ["2-8", "2-10", "2-20"],
        "keywords": ["活性", "非活性", "参照モード", "更新モード", "操作不可", "コントロール状態"],
    },
    {
        "nos": ["2-11"],
        "keywords": ["format", "表示状態", "編集状態", "書式"],
    },
    {
        "nos": ["2-12"],
        "keywords": ["必須", "未入力", "必須項目"],
    },
    {
        "nos": ["2-13"],
        "keywords": ["半角英数字", "日付", "数値以外", "英数", "日付項目", "半角空白"],
    },
    {
        "nos": ["2-14"],
        "keywords": ["桁数", "最大桁数", "最大入力桁数", "最小最大数", "文字数", "バイト数", "maxlength"],
    },
    {
        "nos": ["2-15"],
        "keywords": ["境界値", "最大桁数超過", "最大桁数ちょうど", "最大桁数以内", "範囲", "端数"],
    },
    {
        "nos": ["2-18"],
        "keywords": ["0件", "０件", "null", "NULL", "抽出データ", "処理対象データ"],
    },
    {
        "nos": ["2-19"],
        "keywords": ["tab", "Tab", "Shift+Tab", "フォーカス順"],
    },
    {
        "nos": ["2-21", "2-22", "2-32"],
        "keywords": ["エラーメッセージ", "メッセージ", "フォーカスアウト", "ボタン押下時"],
    },
    {
        "nos": ["2-23", "2-25"],
        "keywords": ["ソート", "条件保持", "検索後", "昇順", "降順"],
    },
    {
        "nos": ["2-24"],
        "keywords": ["検索結果", "検索ダイアログ", "検索"],
    },
    {
        "nos": ["2-26"],
        "keywords": ["一覧", "明細一覧", "事象一覧", "帳票一覧"],
    },
    {
        "nos": ["2-28"],
        "keywords": ["画面遷移", "ダイアログ", "起動する", "遷移先"],
    },
    {
        "nos": ["2-29"],
        "keywords": ["db", "DB", "反映", "登録結果", "更新結果"],
    },
]


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _carry_forward(current: str, previous: str) -> str:
    return current if current else previous


def load_checklist_items(
    workbook_path: str | Path,
    sheet_name: str = CHECKLIST_SHEET_NAME,
    start_row: int = CHECKLIST_START_ROW,
    end_row: int = CHECKLIST_END_ROW,
) -> List[Dict[str, Any]]:
    """
    读取 checklist Excel 的指定 sheet / 行范围，返回结构化项目列表。

    预期读取:
    - sheet: UT仕様書チェックシート
    - rows : 26 ～ 116

    输出字段:
    - no: Checklist No.（例: 2-26）
    - category_l1: 大分类（例: 画面）
    - category_l2: 中分类（例: 初期値 / I/O / 一覧）
    - category_l3: 预留字段（当前先不重点使用）
    - text: 检查内容
    - note: 备注 / 障害事例
    - enabled: 是否打了「〇」
    - row_index: Excel 行号
    """
    workbook = load_workbook(filename=workbook_path, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}")

    ws = workbook[sheet_name]

    items: List[Dict[str, Any]] = []

    prev_category_l1 = ""
    prev_category_l2 = ""
    prev_category_l3 = ""

    for row in range(start_row, end_row + 1):
        # 实际表结构（UT仕様書チェックシート 26～116行）:
        # A列: 章节/空白
        # B列: Checklist No.（例: 2-1）
        # C列: 辅助列（多数情况下为空，先不重点使用）
        # D列: 大分类（例: 画面）
        # E列: 中分类（例: 画面構成 / 初期値 / I/O / 一覧）
        # F列: 检查内容
        # G列: 备注
        # H列: 适用标记(〇)
        no = _cell_text(ws.cell(row=row, column=2).value)
        category_l1 = _cell_text(ws.cell(row=row, column=4).value)
        category_l2 = _cell_text(ws.cell(row=row, column=5).value)
        category_l3 = _cell_text(ws.cell(row=row, column=3).value)
        text = _cell_text(ws.cell(row=row, column=6).value)
        note = _cell_text(ws.cell(row=row, column=7).value)
        enabled_mark = _cell_text(ws.cell(row=row, column=8).value)

        # 合并单元格导致的空值，沿用上一行分类
        category_l1 = _carry_forward(category_l1, prev_category_l1)
        category_l2 = _carry_forward(category_l2, prev_category_l2)
        category_l3 = _carry_forward(category_l3, prev_category_l3)

        prev_category_l1 = category_l1
        prev_category_l2 = category_l2
        prev_category_l3 = category_l3

        # Checklist No. 为空的行不作为有效 checklist 项读取
        if not no:
            continue

        items.append(
            {
                "row_index": row,
                "no": no,
                "category_l1": category_l1,
                "category_l2": category_l2,
                "category_l3": category_l3,
                "text": text,
                "note": note,
                "enabled": enabled_mark == "〇",
            }
        )

    return items

def save_checklist_json(items: List[Dict[str, Any]], output_path: str | Path) -> None:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(items, f, ensure_ascii=False, indent=2)

def find_checklist_item_by_no(
    items: List[Dict[str, Any]],
    checklist_no: str,
) -> Optional[Dict[str, Any]]:
    target = str(checklist_no).strip()
    for item in items:
        if str(item.get("no", "")).strip() == target:
            return item
    return None


# --- Matching/mapping helpers ---

def _normalize_for_match(text: str) -> str:
    normalized = str(text or "").strip().lower()
    normalized = normalized.replace("　", " ")
    normalized = normalized.replace("（", "(").replace("）", ")")
    normalized = normalized.replace("：", ":")
    normalized = re.sub(r"\s+", " ", normalized)
    return normalized


def _find_matching_nos_from_rules(search_text: str) -> List[str]:
    normalized = _normalize_for_match(search_text)
    matched: List[str] = []

    for rule in CHECKLIST_MATCH_RULES:
        keywords = rule.get("keywords", [])
        if any(_normalize_for_match(keyword) in normalized for keyword in keywords):
            for no in rule.get("nos", []):
                if no not in matched:
                    matched.append(no)

    return matched


def map_confirmation_to_checklist_nos(
    confirmation_text: str,
    testcase_category: str = "",
    viewpoint_title: str = "",
    source_basis: Optional[List[str]] = None,
) -> List[str]:
    search_parts = [
        str(confirmation_text or ""),
        str(testcase_category or ""),
        str(viewpoint_title or ""),
        " ".join([str(value) for value in (source_basis or [])]),
    ]
    search_text = " ".join([part for part in search_parts if part]).strip()
    matched = _find_matching_nos_from_rules(search_text)

    # 进一步补强一些组合规则
    normalized = _normalize_for_match(search_text)

    if "フォーカスアウト" in search_text and "2-22" not in matched:
        matched.append("2-22")

    if ("最大桁数" in search_text or "maxlength" in normalized) and "2-14" not in matched:
        matched.append("2-14")

    if ("最大桁数超過" in search_text or "境界" in search_text) and "2-15" not in matched:
        matched.append("2-15")

    if ("一覧" in search_text or "明細一覧" in search_text or "事象一覧" in search_text or "帳票一覧" in search_text) and "2-26" not in matched:
        matched.append("2-26")

    if ("0件" in search_text or "０件" in search_text or "NULL" in search_text or "null" in normalized) and "2-18" not in matched:
        matched.append("2-18")

    if "ソート" in search_text and "2-25" not in matched:
        matched.append("2-25")

    if "検索" in search_text and "2-24" not in matched:
        matched.append("2-24")

    if "画面遷移" in search_text or "ダイアログ" in search_text:
        if "2-28" not in matched:
            matched.append("2-28")

    if "DB" in search_text or "db" in normalized or "反映" in search_text:
        if "2-29" not in matched:
            matched.append("2-29")

    return matched


def attach_checklist_nos_to_confirmations(testcases_data: Dict[str, Any]) -> Dict[str, Any]:
    definitions = testcases_data.get("definitions", {})
    confirmations = definitions.get("confirmations", [])
    viewpoint_map = {
        str(vp.get("id", "")).strip(): vp
        for vp in testcases_data.get("viewpoints", [])
        if str(vp.get("id", "")).strip()
    }

    confirmation_to_cases: Dict[str, List[Dict[str, Any]]] = {}
    for test_case in testcases_data.get("test_cases", []):
        for cid in test_case.get("confirmation_ids", []):
            confirmation_to_cases.setdefault(str(cid).strip(), []).append(test_case)

    for item in confirmations:
        cid = str(item.get("id", "")).strip()
        text = str(item.get("text", "")).strip()
        matched_nos: List[str] = []

        related_cases = confirmation_to_cases.get(cid, [])
        for test_case in related_cases:
            viewpoint_id = str(test_case.get("test_viewpoint_id", "")).strip()
            viewpoint = viewpoint_map.get(viewpoint_id, {})
            case_matches = map_confirmation_to_checklist_nos(
                confirmation_text=text,
                testcase_category=str(test_case.get("category", "")),
                viewpoint_title=str(test_case.get("test_viewpoint", "")) or str(viewpoint.get("title", "")),
                source_basis=test_case.get("source_basis", []) or viewpoint.get("source_basis", []),
            )
            for no in case_matches:
                if no not in matched_nos:
                    matched_nos.append(no)

        if not matched_nos:
            matched_nos = map_confirmation_to_checklist_nos(confirmation_text=text)

        item["checklist_nos"] = matched_nos

    return testcases_data


def main() -> None:
    sample_path = Path("sample_input/BR_チェックリスト.xlsx")
    output_path = Path("output/checklist.json")

    if not sample_path.exists():
        print(f"Checklist file not found: {sample_path}")
        return

    items = load_checklist_items(sample_path)

    save_checklist_json(items, output_path)

    print(f"checklist.json generated: {output_path}")
    print(f"total items: {len(items)}")

    # debug sample
    sample_text = "債権明細一覧が0件で表示されること"
    print("sample mapping:", sample_text, "->", map_confirmation_to_checklist_nos(sample_text))


if __name__ == "__main__":
    main()