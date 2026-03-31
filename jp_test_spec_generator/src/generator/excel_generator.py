from __future__ import annotations

import copy
import json
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


THIN_SIDE = Side(style="thin", color="000000")
HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
SECTION_FILL = PatternFill("solid", fgColor="BDD7EE")
SUBHEADER_FILL = PatternFill("solid", fgColor="E2F0D9")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
CIRCLE_FONT = Font(name="ＭＳ ゴシック")

# PCL colors
PCL_COLOR_MAP = {
    "N": PatternFill("solid", fgColor="C6EFCE"),  # green
    "E": PatternFill("solid", fgColor="FFC7CE"),  # red
    "L": PatternFill("solid", fgColor="FFEB9C"),  # yellow
    "I": PatternFill("solid", fgColor="BDD7EE"),  # blue
}

# Missing checklist highlight (yellow)
MISSING_CHECKLIST_FILL = PatternFill("solid", fgColor="FFFF00")

# Template layout
CHECK_START_ROW = 4
CHECK_END_ROW = 33
ACTION_START_ROW = 34
ACTION_END_ROW = 66
CONFIRM_START_ROW = 67
CONFIRM_END_ROW = 158
PCL_ROW = 161
CASE_HEADER_ROW = 3
CASE_START_COL = 10  # J列
CASE_TEMPLATE_COPY_COL = 45  # AS列
CHECKLIST_NO_COL = 9  # I列
ITEM_NO_COL = 2      # B列
ITEM_TEXT_COL = 3    # C列
TEMPLATE_ITEM_TEXT_COL = 2  # テンプレート主表では B列（結合セルの左上）に説明文を書く
SECTION_LABEL_COL = 1  # A列
CHECK_TEMPLATE_COPY_ROW = 32
ACTION_TEMPLATE_COPY_ROW = 65
CONFIRM_TEMPLATE_COPY_ROW = 157
TEMPLATE_PATH = Path(__file__).resolve().parents[2] / "Sample_input" / "単体テスト仕様書base.xlsx"



class ExcelGenerationError(Exception):
    pass


def _print_progress(message: str) -> None:
    sys.stdout.write(f"{message}\n")
    sys.stdout.flush()


def _print_case_progress(prefix: str, current: int, total: int) -> None:
    if total <= 0:
        print(f"{prefix}: 0/0", end="\r", flush=True)
        return

    percent = int((current / total) * 100)
    print(f"{prefix}: {current}/{total} ({percent}%)", end="\r", flush=True)
    if current >= total:
        print("", flush=True)



def _load_json(json_path: str | Path) -> Dict[str, Any]:
    path = Path(json_path)
    if not path.exists():
        raise ExcelGenerationError(f"JSON file not found: {path}")

    with path.open("r", encoding="utf-8") as f:
        data = json.load(f)

    if not isinstance(data, dict):
        raise ExcelGenerationError("Invalid JSON structure: root must be an object")

    return data



def _validate_testcases(data: Dict[str, Any]) -> None:
    if "definitions" not in data:
        raise ExcelGenerationError("Invalid JSON structure: 'definitions' is required")
    if "test_cases" not in data:
        raise ExcelGenerationError("Invalid JSON structure: 'test_cases' is required")

    definitions = data["definitions"]
    required_definition_keys = ["check_conditions", "actions", "confirmations"]
    for key in required_definition_keys:
        if key not in definitions:
            raise ExcelGenerationError(f"Invalid JSON structure: 'definitions.{key}' is required")

def _is_delivery_excluded_confirmation_text(text: str) -> bool:
    normalized = str(text or "").strip()
    if not normalized:
        return False

    excluded_patterns = [
        "maxlength=",
        "フォーカスアウト時の桁数チェック対象",
        "対象コントロール数:",
    ]
    return any(pattern in normalized for pattern in excluded_patterns)


def _prepare_delivery_view_data(data: Dict[str, Any]) -> Dict[str, Any]:
    prepared = copy.deepcopy(data)
    definitions = prepared.get("definitions", {})
    confirmations = definitions.get("confirmations", [])

    excluded_confirmation_ids = {
        str(item.get("id", "")).strip()
        for item in confirmations
        if _is_delivery_excluded_confirmation_text(str(item.get("text", "")))
    }

    definitions["confirmations"] = [
        item
        for item in confirmations
        if str(item.get("id", "")).strip() not in excluded_confirmation_ids
    ]

    for test_case in prepared.get("test_cases", []):
        test_case["confirmation_ids"] = [
            cid
            for cid in test_case.get("confirmation_ids", [])
            if str(cid).strip() not in excluded_confirmation_ids
        ]

    summary = prepared.get("summary", {})
    if isinstance(summary, dict):
        summary["delivery_excluded_confirmation_count"] = len(excluded_confirmation_ids)
        summary["confirmation_count"] = len(definitions.get("confirmations", []))

    return prepared


# Helper to format checklist numbers for output in Excel
def _format_checklist_nos(item: Dict[str, Any]) -> str:
    checklist_nos = item.get("checklist_nos", []) or []
    values = [str(value).strip() for value in checklist_nos if str(value).strip()]
    if not values:
        return "2-1"
    return ",".join(values)

def _apply_border_range(ws, start_row: int, end_row: int, start_col: int, end_col: int) -> None:
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            ws.cell(row=row, column=col).border = Border(
                left=THIN_SIDE,
                right=THIN_SIDE,
                top=THIN_SIDE,
                bottom=THIN_SIDE,
            )



def _copy_cell_style(src_cell, dst_cell) -> None:
    if src_cell.__class__.__name__ == "MergedCell" or dst_cell.__class__.__name__ == "MergedCell":
        return

    if src_cell.has_style:
        dst_cell._style = copy.copy(src_cell._style)
    if src_cell.font:
        dst_cell.font = copy.copy(src_cell.font)
    if src_cell.fill:
        dst_cell.fill = copy.copy(src_cell.fill)
    if src_cell.border:
        dst_cell.border = copy.copy(src_cell.border)
    if src_cell.alignment:
        dst_cell.alignment = copy.copy(src_cell.alignment)
    if src_cell.number_format:
        dst_cell.number_format = src_cell.number_format
    if src_cell.protection:
        dst_cell.protection = copy.copy(src_cell.protection)



def _copy_row_style(ws: Worksheet, source_row: int, target_row: int, max_col: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, max_col + 1):
        _copy_cell_style(ws.cell(source_row, col), ws.cell(target_row, col))


# Helper to copy merged-cell structure for a single row.
def _copy_row_merges(ws: Worksheet, source_row: int, target_row: int) -> None:
    existing_ranges = {str(rng) for rng in ws.merged_cells.ranges}
    for merged_range in list(ws.merged_cells.ranges):
        if merged_range.min_row == source_row and merged_range.max_row == source_row:
            new_range = f"{get_column_letter(merged_range.min_col)}{target_row}:{get_column_letter(merged_range.max_col)}{target_row}"
            if new_range not in existing_ranges:
                ws.merge_cells(new_range)
                existing_ranges.add(new_range)




def _copy_column_style(ws: Worksheet, source_col: int, target_col: int, max_row: int) -> None:
    source_letter = get_column_letter(source_col)
    target_letter = get_column_letter(target_col)
    case_base_letter = get_column_letter(CASE_START_COL)
    ws.column_dimensions[target_letter].width = ws.column_dimensions[case_base_letter].width
    ws.column_dimensions[target_letter].hidden = ws.column_dimensions[source_letter].hidden
    for row in range(1, max_row + 1):
        _copy_cell_style(ws.cell(row, source_col), ws.cell(row, target_col))


# Helper to adjust row height for wrapped text in column B/C.
def _adjust_wrapped_row_height(ws: Worksheet, row: int, text: str, base_height: float = 18.0) -> None:
    content = str(text or "")
    if not content:
        return

    # B列の結合セル表示を前提に、おおよその文字数で行数を算出する
    approx_chars_per_line = 28
    line_count = max(1, (len(content) // approx_chars_per_line) + (1 if len(content) % approx_chars_per_line else 0))

    # テンプレート行高を基準にする
    template_height = ws.row_dimensions[row].height or base_height
    ws.row_dimensions[row].height = template_height * line_count




def _clear_cell_value_keep_style(ws: Worksheet, row: int, col: int) -> None:
    cell = ws.cell(row=row, column=col)
    if cell.__class__.__name__ == "MergedCell":
        return
    cell.value = None


# Helper to safely set cell value, handling merged cells.
def _set_cell_value_safe(ws: Worksheet, row: int, col: int, value: Any) -> None:
    cell = ws.cell(row=row, column=col)
    if cell.__class__.__name__ != "MergedCell":
        cell.value = value
        return

    for merged_range in ws.merged_cells.ranges:
        if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= col <= merged_range.max_col:
            # Merged child cells cannot be written directly.
            # When clearing a non-anchor merged cell, just skip.
            if value is None and (row != merged_range.min_row or col != merged_range.min_col):
                return

            anchor_cell = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
            if anchor_cell.__class__.__name__ == "MergedCell":
                return

            anchor_cell.value = value
            return

    # Fallback: if this is still a merged child cell, do nothing.
    if cell.__class__.__name__ == "MergedCell":
        return
    cell.value = value



def _find_case_template_capacity(ws: Worksheet) -> int:
    col = CASE_START_COL
    count = 0
    while True:
        value = ws.cell(row=CASE_HEADER_ROW, column=col).value
        if value is None or str(value).strip() == "":
            break
        count += 1
        col += 1
    return count



def _ensure_case_columns(ws: Worksheet, required_case_count: int) -> None:
    existing_case_count = _find_case_template_capacity(ws)

    if existing_case_count <= 0:
        raise ExcelGenerationError("テンプレートのケース列が見つかりません。")

    if required_case_count <= existing_case_count:
        return

    additional = required_case_count - existing_case_count
    _print_progress(f"ケース列を追加します: 追加 {additional} 列")

    # 追加列の列幅はテンプレートのケース列（J列）に合わせる
    base_case_letter = get_column_letter(CASE_START_COL)
    base_case_width = ws.column_dimensions[base_case_letter].width

    for i in range(additional):
        new_col = ws.max_column + 1
        new_letter = get_column_letter(new_col)
        ws.column_dimensions[new_letter].width = base_case_width

        for row in range(1, ws.max_row + 1):
            _clear_cell_value_keep_style(ws, row, new_col)



def _ensure_section_rows(
    ws: Worksheet,
    section_start_row: int,
    reserved_end_row: int,
    required_count: int,
    template_copy_row: int,
) -> int:
    reserved_count = reserved_end_row - section_start_row + 1
    extra = max(0, required_count - reserved_count)
    if extra > 0:
        _print_progress(f"行を拡張します: 開始行={section_start_row}, 追加 {extra} 行")
    if extra <= 0:
        return 0

    max_col = max(ws.max_column, CASE_START_COL)

    for offset in range(extra):
        source_row = template_copy_row + offset
        insert_at = source_row + 1
        ws.insert_rows(insert_at, 1)
        _copy_row_style(ws, source_row, insert_at, max_col)
        _copy_row_merges(ws, source_row, insert_at)
        for col in range(1, max_col + 1):
            _clear_cell_value_keep_style(ws, insert_at, col)
        _print_case_progress("行作成進捗", offset + 1, extra)

    return extra




def _fill_section_rows(
    ws: Worksheet,
    start_row: int,
    items: List[Dict[str, str]],
    section_label: str,
    relation_key: str,
    test_cases: List[Dict[str, Any]],
) -> None:
    if not items:
        return

    _print_progress(f"{section_label} を書き込みます: {len(items)} 行")

    for index, item in enumerate(items):
        row = start_row + index
        item_id = str(item.get("id", "")).strip()
        item_text = str(item.get("text", "")).strip()

        if index == 0:
            _set_cell_value_safe(ws, row, SECTION_LABEL_COL, section_label)
        else:
            _set_cell_value_safe(ws, row, SECTION_LABEL_COL, None)

        _set_cell_value_safe(ws, row, ITEM_NO_COL, item_id)
        _set_cell_value_safe(ws, row, ITEM_TEXT_COL, item_text)
        ws.cell(row=row, column=ITEM_TEXT_COL).alignment = LEFT
        _adjust_wrapped_row_height(ws, row, item_text)

        is_confirmation = section_label == "確認内容"
        checklist_nos = item.get("checklist_nos", []) or []
        has_values = any(str(v).strip() for v in checklist_nos)

        checklist_text = _format_checklist_nos(item) if is_confirmation else ""
        _set_cell_value_safe(ws, row, CHECKLIST_NO_COL, checklist_text if checklist_text else None)

        checklist_cell = ws.cell(row=row, column=CHECKLIST_NO_COL)
        if checklist_cell.__class__.__name__ != "MergedCell":
            # 取消自动换行
            checklist_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)

            # 未填则标黄
            if is_confirmation and not has_values:
                checklist_cell.fill = MISSING_CHECKLIST_FILL

        for case_index, test_case in enumerate(test_cases):
            col = CASE_START_COL + case_index
            related_ids = test_case.get(relation_key, [])
            circle_value = "○" if item_id in related_ids else None
            _set_cell_value_safe(ws, row, col, circle_value)
            target_cell = ws.cell(row=row, column=col)
            if target_cell.__class__.__name__ == "MergedCell":
                for merged_range in ws.merged_cells.ranges:
                    if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= col <= merged_range.max_col:
                        target_cell = ws.cell(merged_range.min_row, merged_range.min_col)
                        break
            if circle_value == "○":
                target_cell.font = CIRCLE_FONT
        _print_case_progress(f"{section_label} 書き込み進捗", index + 1, len(items))


# テンプレート主表用の行書き込みヘルパ
def _fill_template_section_rows(
    ws: Worksheet,
    start_row: int,
    items: List[Dict[str, str]],
    section_label: str,
    relation_key: str,
    test_cases: List[Dict[str, Any]],
) -> None:
    if not items:
        return

    _print_progress(f"{section_label} をテンプレートへ書き込みます: {len(items)} 行")

    for index, item in enumerate(items):
        row = start_row + index
        item_id = str(item.get("id", "")).strip()
        item_text = str(item.get("text", "")).strip()

        if index == 0:
            _set_cell_value_safe(ws, row, SECTION_LABEL_COL, section_label)
        else:
            _set_cell_value_safe(ws, row, SECTION_LABEL_COL, None)

        _clear_cell_value_keep_style(ws, row, ITEM_NO_COL)

        # テンプレート主表では定義IDは表示せず、説明文だけを結合セル左上へ出力する
        _set_cell_value_safe(ws, row, TEMPLATE_ITEM_TEXT_COL, item_text)
        if ITEM_TEXT_COL != TEMPLATE_ITEM_TEXT_COL:
            _clear_cell_value_keep_style(ws, row, ITEM_TEXT_COL)

        is_confirmation = section_label == "確認内容"
        checklist_nos = item.get("checklist_nos", []) or []
        has_values = any(str(v).strip() for v in checklist_nos)

        checklist_text = _format_checklist_nos(item) if is_confirmation else ""
        _set_cell_value_safe(ws, row, CHECKLIST_NO_COL, checklist_text if checklist_text else None)

        checklist_cell = ws.cell(row=row, column=CHECKLIST_NO_COL)
        if checklist_cell.__class__.__name__ != "MergedCell":
            # 取消自动换行
            checklist_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)

            # 未填则标黄
            if is_confirmation and not has_values:
                checklist_cell.fill = MISSING_CHECKLIST_FILL

        target_text_cell = ws.cell(row=row, column=TEMPLATE_ITEM_TEXT_COL)
        if target_text_cell.__class__.__name__ == "MergedCell":
            for merged_range in ws.merged_cells.ranges:
                if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= TEMPLATE_ITEM_TEXT_COL <= merged_range.max_col:
                    target_text_cell = ws.cell(merged_range.min_row, merged_range.min_col)
                    break
        target_text_cell.alignment = LEFT
        _adjust_wrapped_row_height(ws, row, item_text)

        for case_index, test_case in enumerate(test_cases):
            col = CASE_START_COL + case_index
            related_ids = test_case.get(relation_key, [])
            circle_value = "○" if item_id in related_ids else None
            _set_cell_value_safe(ws, row, col, circle_value)
            target_cell = ws.cell(row=row, column=col)
            if target_cell.__class__.__name__ == "MergedCell":
                for merged_range in ws.merged_cells.ranges:
                    if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= col <= merged_range.max_col:
                        target_cell = ws.cell(merged_range.min_row, merged_range.min_col)
                        break
            if circle_value == "○":
                target_cell.font = CIRCLE_FONT

        _print_case_progress(f"{section_label} テンプレート書き込み進捗", index + 1, len(items))



def _clear_unused_rows(ws: Worksheet, start_row: int, reserved_count: int, used_count: int, max_col: int) -> None:
    for row in range(start_row + used_count, start_row + reserved_count):
        for col in range(SECTION_LABEL_COL, max_col + 1):
            _clear_cell_value_keep_style(ws, row, col)
        _clear_cell_value_keep_style(ws, row, TEMPLATE_ITEM_TEXT_COL)
        if ITEM_TEXT_COL != TEMPLATE_ITEM_TEXT_COL:
            _clear_cell_value_keep_style(ws, row, ITEM_TEXT_COL)



def _fill_case_header(ws: Worksheet, test_cases: List[Dict[str, Any]], pcl_row: int) -> None:
    _print_progress(f"ケース列ヘッダとPCLを書き込みます: {len(test_cases)} 件")
    for case_index, test_case in enumerate(test_cases):
        col = CASE_START_COL + case_index
        _set_cell_value_safe(ws, CASE_HEADER_ROW, col, test_case.get("case_no") or test_case.get("case_id"))
        header_cell = ws.cell(row=CASE_HEADER_ROW, column=col)
        if header_cell.__class__.__name__ == "MergedCell":
            for merged_range in ws.merged_cells.ranges:
                if merged_range.min_row <= CASE_HEADER_ROW <= merged_range.max_row and merged_range.min_col <= col <= merged_range.max_col:
                    header_cell = ws.cell(merged_range.min_row, merged_range.min_col)
                    break
        header_cell.alignment = CENTER

        pcl_list = test_case.get("pcl", [])
        pcl_text = pcl_list[0] if pcl_list else "N"
        _set_cell_value_safe(ws, pcl_row, col, pcl_text)
        pcl_cell = ws.cell(pcl_row, col)
        if pcl_cell.__class__.__name__ == "MergedCell":
            for merged_range in ws.merged_cells.ranges:
                if merged_range.min_row <= pcl_row <= merged_range.max_row and merged_range.min_col <= col <= merged_range.max_col:
                    pcl_cell = ws.cell(merged_range.min_row, merged_range.min_col)
                    break
        pcl_cell.alignment = CENTER
        fill = PCL_COLOR_MAP.get(pcl_text)
        if fill:
            pcl_cell.fill = fill
        _print_case_progress("ケースヘッダ書き込み進捗", case_index + 1, len(test_cases))



def _clear_unused_case_columns(ws: Worksheet, start_col: int, used_count: int, max_clear_row: int) -> None:
    col = start_col + used_count
    while col <= ws.max_column:
        header_value = ws.cell(row=CASE_HEADER_ROW, column=col).value
        if header_value is None or str(header_value).strip() == "":
            break
        for row in range(1, max_clear_row + 1):
            _clear_cell_value_keep_style(ws, row, col)
        col += 1



def _write_metadata(ws: Worksheet, screen_id: str, screen_name: str) -> None:
    _set_cell_value_safe(ws, 1, 2, screen_id)
    _set_cell_value_safe(ws, 2, 2, screen_name)



def _write_case_detail_sheet(wb: Workbook, test_cases: List[Dict[str, Any]]) -> None:
    if "CaseList" in wb.sheetnames:
        del wb["CaseList"]

    ws = wb.create_sheet("CaseList")
    headers = [
        "case_no",
        "case_id",
        "category",
        "test_viewpoint_id",
        "test_viewpoint",
        "pcl",
        "check_condition_ids",
        "action_ids",
        "confirmation_ids",
        "source_basis",
    ]

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.font = Font(bold=True)
        cell.alignment = CENTER
        cell.fill = HEADER_FILL

    for row_index, test_case in enumerate(test_cases, start=2):
        ws.cell(row=row_index, column=1).value = test_case.get("case_no")
        ws.cell(row=row_index, column=2).value = test_case.get("case_id")
        ws.cell(row=row_index, column=3).value = test_case.get("category")
        ws.cell(row=row_index, column=4).value = test_case.get("test_viewpoint_id")
        ws.cell(row=row_index, column=5).value = test_case.get("test_viewpoint")
        ws.cell(row=row_index, column=6).value = ", ".join(test_case.get("pcl", []))
        ws.cell(row=row_index, column=7).value = ", ".join(test_case.get("check_condition_ids", []))
        ws.cell(row=row_index, column=8).value = ", ".join(test_case.get("action_ids", []))        
        # confirmation_ids は delivery view 用に事前フィルタ済み
        ws.cell(row=row_index, column=9).value = ", ".join(test_case.get("confirmation_ids", []))
        ws.cell(row=row_index, column=10).value = ", ".join(test_case.get("source_basis", []))

    widths = {
        "A": 12,
        "B": 12,
        "C": 14,
        "D": 16,
        "E": 42,
        "F": 10,
        "G": 20,
        "H": 20,
        "I": 20,
        "J": 30,
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    _apply_border_range(ws, 1, max(2, len(test_cases) + 1), 1, len(headers))



def _write_definitions_sheet(
    wb: Workbook,
    check_conditions: List[Dict[str, str]],
    actions: List[Dict[str, str]],
    confirmations: List[Dict[str, str]],
) -> None:
    if "Definitions" in wb.sheetnames:
        del wb["Definitions"]

    ws = wb.create_sheet("Definitions")
    headers = ["type", "id", "text"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.font = Font(bold=True)
        cell.alignment = CENTER
        cell.fill = HEADER_FILL

    definition_row = 2
    for section_type, section_items in [
        ("check_condition", check_conditions),
        ("action", actions),
        ("confirmation", confirmations),
    ]:
        for item in section_items:
            ws.cell(row=definition_row, column=1).value = section_type
            ws.cell(row=definition_row, column=2).value = item.get("id")
            ws.cell(row=definition_row, column=3).value = item.get("text")
            definition_row += 1

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 60
    _apply_border_range(ws, 1, max(2, definition_row - 1), 1, 3)



def _generate_from_template(
    data: Dict[str, Any],
    output_path: str | Path,
    template_path: str | Path,
    sheet_name: Optional[str] = None,
) -> Path:
    template = Path(template_path)
    if not template.exists():
        raise ExcelGenerationError(f"Template file not found: {template}")

    wb = load_workbook(template)
    _print_progress(f"テンプレートを読み込みました: {template}")
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    prepared_data = _prepare_delivery_view_data(data)

    screen_id = str(prepared_data.get("screen_id", "")).strip()
    screen_name = str(prepared_data.get("screen_name", "")).strip()
    definitions = prepared_data["definitions"]
    test_cases = prepared_data.get("test_cases", [])
    check_conditions = definitions.get("check_conditions", [])
    actions = definitions.get("actions", [])
    confirmations = definitions.get("confirmations", [])

    _print_progress(
        f"生成開始: ケース数={len(test_cases)}, チェック条件={len(check_conditions)}, アクション={len(actions)}, 確認内容={len(confirmations)}"
    )

    _ensure_case_columns(ws, len(test_cases))

    check_extra = _ensure_section_rows(
        ws,
        CHECK_START_ROW,
        CHECK_END_ROW,
        len(check_conditions),
        CHECK_TEMPLATE_COPY_ROW,
    )
    action_start = ACTION_START_ROW + check_extra
    action_end = ACTION_END_ROW + check_extra

    action_extra = _ensure_section_rows(
        ws,
        action_start,
        action_end,
        len(actions),
        ACTION_TEMPLATE_COPY_ROW + check_extra,
    )
    confirm_start = CONFIRM_START_ROW + check_extra + action_extra
    confirm_end = CONFIRM_END_ROW + check_extra + action_extra

    confirm_extra = _ensure_section_rows(
        ws,
        confirm_start,
        confirm_end,
        len(confirmations),
        CONFIRM_TEMPLATE_COPY_ROW + check_extra + action_extra,
    )
    pcl_row = PCL_ROW + check_extra + action_extra + confirm_extra

    _print_progress("テンプレートへの書き込みを開始します")
    _write_metadata(ws, screen_id, screen_name)
    _fill_case_header(ws, test_cases, pcl_row)

    _fill_template_section_rows(ws, CHECK_START_ROW, check_conditions, "チェック条件", "check_condition_ids", test_cases)
    _fill_template_section_rows(ws, action_start, actions, "アクション", "action_ids", test_cases)
    _fill_template_section_rows(ws, confirm_start, confirmations, "確認内容", "confirmation_ids", test_cases)

    _clear_unused_rows(ws, CHECK_START_ROW, max(0, CHECK_END_ROW - CHECK_START_ROW + 1), len(check_conditions), ws.max_column)
    _clear_unused_rows(ws, action_start, max(0, action_end - action_start + 1), len(actions), ws.max_column)
    _clear_unused_rows(ws, confirm_start, max(0, confirm_end - confirm_start + 1), len(confirmations), ws.max_column)
    _clear_unused_case_columns(ws, CASE_START_COL, len(test_cases), pcl_row)

    _set_cell_value_safe(ws, pcl_row, SECTION_LABEL_COL, "PCL区分")
    ws.freeze_panes = f"{get_column_letter(CASE_START_COL)}{CHECK_START_ROW}"

    _print_progress("補助シートを書き込みます")
    _write_definitions_sheet(wb, check_conditions, actions, confirmations)
    _write_case_detail_sheet(wb, test_cases)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    _print_progress(f"Excelを保存します: {output}")
    wb.save(output)
    _print_progress("Excel保存が完了しました")
    return output



def _generate_simple_matrix(
    data: Dict[str, Any],
    output_path: str | Path,
    sheet_name: str = "TestSpec",
) -> Path:
    prepared_data = _prepare_delivery_view_data(data)

    screen_id = str(prepared_data.get("screen_id", "")).strip()
    screen_name = str(prepared_data.get("screen_name", "")).strip()
    definitions = prepared_data["definitions"]
    test_cases = prepared_data.get("test_cases", [])

    check_conditions = definitions.get("check_conditions", [])
    actions = definitions.get("actions", [])
    confirmations = definitions.get("confirmations", [])

    _print_progress(
        f"簡易マトリクス生成開始: ケース数={len(test_cases)}, チェック条件={len(check_conditions)}, アクション={len(actions)}, 確認内容={len(confirmations)}"
    )

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.freeze_panes = "E8"
    ws.sheet_view.showGridLines = False

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(5, 4 + len(test_cases)))
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = f"単体テスト仕様書マトリクス - {screen_id} {screen_name}".strip()
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = LEFT
    title_cell.fill = HEADER_FILL

    row = 5
    ws.cell(row=row, column=1).value = "区分"
    ws.cell(row=row, column=2).value = "項番"
    ws.cell(row=row, column=3).value = "内容"
    ws.cell(row=row, column=4).value = "チェックリストNo."
    ws.cell(row=row, column=4).alignment = CENTER
    ws.cell(row=row, column=4).fill = HEADER_FILL
    for case_index, test_case in enumerate(test_cases):
        col = 5 + case_index
        ws.cell(row=row, column=col).value = test_case.get("case_no") or test_case.get("case_id")
        ws.cell(row=row, column=col).alignment = CENTER
        ws.cell(row=row, column=col).fill = HEADER_FILL
    row += 1

    ws.cell(row=row, column=1).value = "PCL"
    for case_index, test_case in enumerate(test_cases):
        col = 5 + case_index
        pcl_value = (test_case.get("pcl") or ["N"])[0]
        ws.cell(row=row, column=col).value = pcl_value
        ws.cell(row=row, column=col).alignment = CENTER
        fill = PCL_COLOR_MAP.get(pcl_value)
        if fill:
            ws.cell(row=row, column=col).fill = fill
    row += 1

    for section_title, items, relation_key in [
        ("チェック条件", check_conditions, "check_condition_ids"),
        ("アクション", actions, "action_ids"),
        ("確認内容", confirmations, "confirmation_ids"),
    ]:
        start_row = row
        end_row = row + max(0, len(items) - 1)
        if items:
            ws.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
            ws.cell(row=start_row, column=1).value = section_title
            ws.cell(row=start_row, column=1).alignment = CENTER
            ws.cell(row=start_row, column=1).fill = SECTION_FILL
        for index, item in enumerate(items):
            current_row = row + index
            item_id = str(item.get("id", "")).strip()
            item_text = str(item.get("text", "")).strip()
            ws.cell(row=current_row, column=2).value = item_id
            ws.cell(row=current_row, column=3).value = item_text
            ws.cell(row=current_row, column=3).alignment = LEFT
            _adjust_wrapped_row_height(ws, current_row, item_text)

            is_confirmation = section_title == "確認内容"
            checklist_nos = item.get("checklist_nos", []) or []
            has_values = any(str(v).strip() for v in checklist_nos)

            checklist_text = _format_checklist_nos(item) if is_confirmation else ""
            cell = ws.cell(row=current_row, column=4)
            cell.value = checklist_text if checklist_text else None

            # 取消自动换行
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)

            # 未填则标黄
            if is_confirmation and not has_values:
                cell.fill = MISSING_CHECKLIST_FILL

            for case_index, test_case in enumerate(test_cases):
                col = 5 + case_index
                related_ids = test_case.get(relation_key, [])
                circle_value = "○" if item_id in related_ids else None
                ws.cell(row=current_row, column=col).value = circle_value
                if circle_value == "○":
                    ws.cell(row=current_row, column=col).font = CIRCLE_FONT
        row += len(items)

    _print_progress("補助シートを書き込みます")
    _write_definitions_sheet(wb, check_conditions, actions, confirmations)
    _write_case_detail_sheet(wb, test_cases)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    _print_progress(f"Excelを保存します: {output}")
    wb.save(output)
    _print_progress("Excel保存が完了しました")
    return output



def generate_excel(
    json_path: str | Path,
    output_path: str | Path,
    sheet_name: str = "TestSpec",
    template_path: Optional[str | Path] = None,
) -> Path:
    data = _load_json(json_path)
    _validate_testcases(data)

    resolved_template = Path(template_path) if template_path else TEMPLATE_PATH
    if resolved_template.exists():
        return _generate_from_template(
            data=data,
            output_path=output_path,
            template_path=resolved_template,
            sheet_name=sheet_name,
        )

    return _generate_simple_matrix(
        data=data,
        output_path=output_path,
        sheet_name=sheet_name,
    )



def main() -> None:
    import argparse

    parser = argparse.ArgumentParser(description="Generate Excel test specification from testcases.json")
    parser.add_argument("json_path", help="Path to testcases.json")
    parser.add_argument("output_path", help="Path to output .xlsx file")
    parser.add_argument("--sheet-name", default="TestSpec", help="Main sheet name")
    parser.add_argument("--template-path", default=str(TEMPLATE_PATH), help="Path to template .xlsx file")
    args = parser.parse_args()

    output = generate_excel(
        json_path=args.json_path,
        output_path=args.output_path,
        sheet_name=args.sheet_name,
        template_path=args.template_path,
    )
    print(f"Excel generated: {output}")


if __name__ == "__main__":
    main()
