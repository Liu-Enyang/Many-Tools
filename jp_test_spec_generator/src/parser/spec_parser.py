from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List

from openpyxl import load_workbook


def parse_spec(file_path: str | Path) -> Dict[str, Any]:
    path = Path(file_path)

    if not path.exists():
        return {
            "file": str(path),
            "sheets": [],
            "raw_summary": {
                "sheet_count": 0,
                "note": "spec file not found",
            },
        }

    wb = load_workbook(path, data_only=True)
    sheets: List[Dict[str, Any]] = []

    for ws in wb.worksheets:
        rows: List[List[Any]] = []
        for row in ws.iter_rows(values_only=True):
            values = [cell for cell in row]
            if any(v is not None and str(v).strip() != "" for v in values):
                rows.append(values)

        sheets.append({
            "sheet_name": ws.title,
            "rows": rows[:200],  # 先限制一点，避免太大
            "row_count": len(rows),
        })

    return {
        "file": str(path),
        "sheets": sheets,
        "raw_summary": {
            "sheet_count": len(sheets),
        },
    }