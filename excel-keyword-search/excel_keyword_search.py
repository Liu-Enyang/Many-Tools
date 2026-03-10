import os
import threading
import tkinter as tk
from tkinter import filedialog, ttk, messagebox
from openpyxl import load_workbook, Workbook
import xlrd


class ExcelKeywordSearchApp:

    def __init__(self, root):

        self.root = root
        self.root.title("Excel Keyword Search Tool")
        self.root.geometry("650x350")

        self.folder_path = tk.StringVar()
        self.keyword = tk.StringVar()

        self.results = []

        self.create_ui()

    def create_ui(self):

        frame = tk.Frame(self.root)
        frame.pack(pady=20)

        tk.Label(frame, text="Folder:").grid(row=0, column=0, sticky="w")

        tk.Entry(frame, textvariable=self.folder_path, width=50).grid(row=0, column=1)

        tk.Button(frame, text="Browse", command=self.select_folder).grid(row=0, column=2, padx=5)

        tk.Label(frame, text="Keyword:").grid(row=1, column=0, pady=10, sticky="w")

        tk.Entry(frame, textvariable=self.keyword, width=50).grid(row=1, column=1)

        self.search_btn = tk.Button(self.root, text="Start Search", height=2, width=20, command=self.start_search)
        self.search_btn.pack(pady=10)

        self.progress = ttk.Progressbar(self.root, length=500)
        self.progress.pack(pady=10)

        self.file_label = tk.Label(self.root, text="Current file: -")
        self.file_label.pack()

        self.result_label = tk.Label(self.root, text="Results found: 0")
        self.result_label.pack()

    def select_folder(self):

        folder = filedialog.askdirectory()

        if folder:
            self.folder_path.set(folder)

    def start_search(self):

        if not self.folder_path.get():
            messagebox.showwarning("Warning", "Please select folder")
            return

        if not self.keyword.get():
            messagebox.showwarning("Warning", "Please enter keyword")
            return

        self.search_btn.config(state="disabled")
        self.results = []

        thread = threading.Thread(target=self.search)
        thread.start()

    def get_excel_files(self, folder):

        excel_files = []

        for root, dirs, files in os.walk(folder):
            for file in files:

                if file.endswith((".xlsx", ".xlsm", ".xls")):
                    excel_files.append(os.path.join(root, file))

        return excel_files

    def search(self):

        folder = self.folder_path.get()
        keyword = self.keyword.get()

        excel_files = self.get_excel_files(folder)

        if not excel_files:
            messagebox.showinfo("Info", "No Excel files found")
            return

        self.progress["maximum"] = len(excel_files)

        for i, file_path in enumerate(excel_files):

            self.file_label.config(text=f"Current file: {os.path.basename(file_path)}")

            try:

                if file_path.endswith(".xls"):
                    self.search_xls(file_path, keyword)
                else:
                    self.search_xlsx(file_path, keyword)

            except Exception:
                pass

            self.progress["value"] = i + 1
            self.result_label.config(text=f"Results found: {len(self.results)}")

        self.save_results()

        self.search_btn.config(state="normal")

    def search_xlsx(self, file_path, keyword):

        wb = load_workbook(file_path, data_only=True)

        for sheet in wb.sheetnames:

            ws = wb[sheet]

            for row in ws.iter_rows():

                for cell in row:

                    if cell.value:

                        value = str(cell.value)

                        if keyword in value:

                            self.results.append([
                                os.path.basename(file_path),
                                sheet,
                                cell.coordinate,
                                value,
                                file_path
                            ])

    def search_xls(self, file_path, keyword):

        wb = xlrd.open_workbook(file_path)

        for sheet in wb.sheets():

            for row in range(sheet.nrows):

                for col in range(sheet.ncols):

                    value = str(sheet.cell_value(row, col))

                    if keyword in value:

                        cell = f"R{row+1}C{col+1}"

                        self.results.append([
                            os.path.basename(file_path),
                            sheet.name,
                            cell,
                            value,
                            file_path
                        ])

    def save_results(self):

        if not self.results:
            messagebox.showinfo("Done", "No results found")
            return

        wb = Workbook()
        ws = wb.active

        ws.append(["File", "Sheet", "Cell", "Value", "Path"])

        for r in self.results:
            ws.append(r)

        save_path = os.path.join(self.folder_path.get(), "search_results.xlsx")

        wb.save(save_path)

        messagebox.showinfo("Done", f"Search finished\nSaved to:\n{save_path}")


if __name__ == "__main__":

    root = tk.Tk()
    app = ExcelKeywordSearchApp(root)
    root.mainloop()