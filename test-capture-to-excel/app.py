import io
import json
import math
import os
import re
import time
import hashlib
import sys
import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from dataclasses import dataclass, field
from datetime import datetime
from typing import List, Optional

from PIL import ImageGrab, Image, ImageDraw
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from pynput import keyboard
try:
    import pystray
except ImportError:
    pystray = None


# =========================
# 设置文件
# =========================
if getattr(sys, 'frozen', False):
    _APP_DIR = os.path.dirname(sys.executable)
else:
    _APP_DIR = os.path.dirname(os.path.abspath(__file__))

_SETTINGS_FILE = os.path.join(_APP_DIR, "settings.json")

_DEFAULT_SETTINGS: dict = {
    "template_path": "",
    "screenshot_marker": "{{screenshot}}",
    "image_scale": 80,   # percentage 1-200
    "excel_zoom": None,  # None = use template default (or 80 for non-template)
    "image_gap": 2,      # blank rows between images
    "excel_title": "【カスタマイズ】　概要画面の対応",
    "sheet_name": "0001",
    "output_dir": "",    # empty = use default (./output)
}


def _load_settings() -> dict:
    if os.path.exists(_SETTINGS_FILE):
        try:
            with open(_SETTINGS_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
            for k, v in _DEFAULT_SETTINGS.items():
                if k not in data:
                    data[k] = v
            return data
        except Exception:
            pass
    return dict(_DEFAULT_SETTINGS)


def _save_settings(settings: dict) -> None:
    try:
        with open(_SETTINGS_FILE, "w", encoding="utf-8") as f:
            json.dump(settings, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"[保存设置失败] {e}")


# =========================
# 数据模型
# =========================
@dataclass
class CaptureItem:
    index: int
    timestamp: datetime
    image_path: str
    hash_value: str


@dataclass
class CaptureSession:
    session_id: str
    session_dir: str
    excel_path: str
    captures: List[CaptureItem] = field(default_factory=list)
    is_recording: bool = False


# =========================
# 核心服务
# =========================
class TestCaptureApp:
    def __init__(self, ui_callback=None) -> None:
        self.base_output_dir = os.path.join(os.getcwd(), "output")
        os.makedirs(self.base_output_dir, exist_ok=True)

        self.session: Optional[CaptureSession] = None
        self.monitor_thread: Optional[threading.Thread] = None
        self.stop_event = threading.Event()

        self.last_saved_hash: Optional[str] = None
        self.last_clipboard_hash: Optional[str] = None

        self.ui_callback = ui_callback

        # Excel 输出配置
        self.excel_title = "【カスタマイズ】　概要画面の対応"
        self.sheet_name = "0001"
        self.output_dir = self.base_output_dir

        # 全局设置
        self.settings = _load_settings()

    # -------------------------
    # UI 通知
    # -------------------------
    def _notify_ui(self) -> None:
        if self.ui_callback:
            self.ui_callback()

    # -------------------------
    # 工具方法
    # -------------------------
    @staticmethod
    def _safe_sheet_name(name: str) -> str:
        invalid_chars = [":", "\\", "/", "?", "*", "[", "]"]
        safe_name = name
        for ch in invalid_chars:
            safe_name = safe_name.replace(ch, "_")
        safe_name = safe_name.strip() or "Sheet1"
        return safe_name[:31]

    @staticmethod
    def _safe_file_name(name: str) -> str:
        safe_name = re.sub(r'[<>:"/\\\\|?*]', "_", name)
        safe_name = safe_name.strip().rstrip(".")
        return safe_name or "sheet"

    @staticmethod
    def _calculate_image_hash(image: Image.Image) -> str:
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        return hashlib.md5(buffer.getvalue()).hexdigest()

    @staticmethod
    def _get_resized_size(width: int, height: int, max_width: int) -> tuple[int, int]:
        if width <= max_width:
            return width, height
        ratio = max_width / width
        return int(width * ratio), int(height * ratio)

    def _create_resized_excel_image(self, image_path: str, scale: float = 0.8) -> XLImage:
        pil_img = Image.open(image_path)
        new_width = int(pil_img.width * scale)
        new_height = int(pil_img.height * scale)
        xl_img = XLImage(image_path)
        xl_img.width = new_width
        xl_img.height = new_height
        return xl_img

    def _get_image_scale(self) -> float:
        """从设置中获取图片缩放比例（返回 0.0~2.0 的小数）"""
        scale = self.settings.get("image_scale", 80)
        try:
            f = float(scale)
            return f / 100.0 if f > 1 else f
        except (TypeError, ValueError):
            return 0.8

    def _get_image_gap(self) -> int:
        """从设置中获取图片间距（空白行数）"""
        gap = self.settings.get("image_gap", 2)
        try:
            return int(gap)
        except (TypeError, ValueError):
            return 2

    def _apply_excel_zoom(self, ws, is_template: bool = False) -> None:
        """应用 Excel 缩放比例；is_template=True 且未设置时保留模板原有缩放"""
        zoom = self.settings.get("excel_zoom", None)
        if zoom is not None:
            try:
                ws.sheet_view.zoomScale = int(zoom)
                return
            except (TypeError, ValueError):
                pass
        if not is_template:
            ws.sheet_view.zoomScale = 80

    def _find_template_markers(self, ws, marker: str) -> list:
        """扫描模板 sheet，找出所有含标记文本的单元格，按行列升序排列"""
        marker_stripped = marker.strip()
        found = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and str(cell.value).strip() == marker_stripped:
                    found.append(cell)
        found.sort(key=lambda c: (c.row, c.column))
        return found

    # -------------------------
    # 会话管理
    # -------------------------
    def start_recording(self) -> bool:
        if self.session and self.session.is_recording:
            print("已经在录制中。")
            return False

        session_id = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_sheet_name_for_file = self._safe_file_name(self.sheet_name)
        session_folder_name = f"{session_id}_{safe_sheet_name_for_file}"
        session_dir = os.path.join(self.output_dir, session_folder_name)
        os.makedirs(session_dir, exist_ok=True)

        excel_path = os.path.join(
            session_dir,
            f"{session_id}_screenshots_{safe_sheet_name_for_file}.xlsx"
        )

        self.session = CaptureSession(
            session_id=session_id,
            session_dir=session_dir,
            excel_path=excel_path,
            captures=[],
            is_recording=True
        )

        self.last_saved_hash = None
        self.last_clipboard_hash = None
        self.stop_event.clear()

        # 记录当前剪贴板避免把旧图保存进去
        try:
            clipboard_obj = ImageGrab.grabclipboard()
            if isinstance(clipboard_obj, Image.Image):
                self.last_clipboard_hash = self._calculate_image_hash(clipboard_obj)
        except Exception:
            pass

        self.monitor_thread = threading.Thread(
            target=self._monitor_clipboard,
            daemon=True
        )
        self.monitor_thread.start()

        print("=" * 60)
        print(f"开始录制: {session_id}")
        print("请使用 Alt + PrintScreen 截取当前活动窗口。")
        print("结束快捷键: Ctrl + Alt + F10")
        print(f"输出目录: {session_dir}")
        print("=" * 60)

        self._notify_ui()
        return True

    def stop_recording(self) -> bool:
        if not self.session or not self.session.is_recording:
            print("当前没有录制中的会话。")
            return False

        self.session.is_recording = False
        self.stop_event.set()

        if self.monitor_thread and self.monitor_thread.is_alive():
            self.monitor_thread.join(timeout=2)

        print("=" * 60)
        print("录制结束，开始导出 Excel...")
        self._export_to_excel()
        print(f"Excel 已生成: {self.session.excel_path}")
        print(f"共保存截图: {len(self.session.captures)} 张")
        print("=" * 60)

        self._notify_ui()
        return True

    def reset_session(self) -> None:
        if self.session and self.session.is_recording:
            return
        self.session = None
        self.last_saved_hash = None
        self.last_clipboard_hash = None
        self._notify_ui()

    # -------------------------
    # 剪贴板监听
    # -------------------------
    def _monitor_clipboard(self) -> None:
        while not self.stop_event.is_set():
            try:
                clipboard_obj = ImageGrab.grabclipboard()

                if clipboard_obj is None:
                    time.sleep(0.3)
                    continue

                if isinstance(clipboard_obj, Image.Image):
                    current_hash = self._calculate_image_hash(clipboard_obj)

                    if current_hash == self.last_clipboard_hash:
                        time.sleep(0.3)
                        continue

                    self.last_clipboard_hash = current_hash
                    self._handle_clipboard_image(clipboard_obj, current_hash)

            except Exception as e:
                print(f"[剪贴板监听异常] {e}")

            time.sleep(0.3)

    def _handle_clipboard_image(self, image: Image.Image, img_hash: str) -> None:
        if not self.session or not self.session.is_recording:
            return

        if img_hash == self.last_saved_hash:
            return

        self.last_saved_hash = img_hash

        index = len(self.session.captures) + 1
        timestamp = datetime.now()
        filename = f"{index:03d}_{timestamp.strftime('%Y%m%d_%H%M%S')}.png"
        image_path = os.path.join(self.session.session_dir, filename)

        image.save(image_path, "PNG")

        item = CaptureItem(
            index=index,
            timestamp=timestamp,
            image_path=image_path,
            hash_value=img_hash
        )
        self.session.captures.append(item)

        print(f"[已保存] 第 {index} 张 -> {image_path}")
        self._notify_ui()

        if self.ui_callback:
            try:
                self.ui_callback("capture_saved", item)
            except TypeError:
                self.ui_callback()

    # -------------------------
    # Excel 写入（无模板模式）
    # -------------------------
    def _write_sheet_content(
        self,
        ws,
        excel_title: str,
        image_paths: List[str],
        image_scale: float = 0.8,
        zoom_scale: int = 80,
        image_gap: int = 2
    ) -> None:
        ws.sheet_view.zoomScale = zoom_scale

        ws["A1"] = excel_title

        current_row = 3
        image_col = "C"
        labels = ["＜前提＞", "＜操作＞", "＜結果＞"]

        for i, image_path in enumerate(image_paths):
            if not os.path.exists(image_path):
                continue

            if i < len(labels):
                ws[f"A{current_row}"] = labels[i]

            img_for_excel = self._create_resized_excel_image(image_path, scale=image_scale)
            ws.add_image(img_for_excel, f"{image_col}{current_row}")

            pil_img = Image.open(image_path)
            display_height = int(pil_img.height * image_scale)
            estimated_rows = max(1, math.ceil(display_height * 0.75 / 20))

            for r in range(current_row, current_row + estimated_rows):
                ws.row_dimensions[r].height = 20

            current_row += estimated_rows + image_gap

    # -------------------------
    # Excel 写入（模板模式）
    # -------------------------
    def _write_template_sheet_content(
        self,
        ws,
        image_paths: List[str],
        image_scale: float,
        marker: str,
        image_gap: int = 2
    ) -> None:
        """找到模板中第一个标记，从该位置起顺序插入所有截图"""
        # 找第一个标记单元格
        first_marker = None
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and str(cell.value).strip() == marker.strip():
                    first_marker = cell
                    break
            if first_marker:
                break

        if first_marker:
            first_marker.value = None  # 清除标记文本
            current_row = first_marker.row
            anchor_col = first_marker.column_letter
        else:
            current_row = 3
            anchor_col = "C"

        for image_path in image_paths:
            if not os.path.exists(image_path):
                continue

            img_for_excel = self._create_resized_excel_image(image_path, scale=image_scale)
            ws.add_image(img_for_excel, f"{anchor_col}{current_row}")

            pil_img = Image.open(image_path)
            display_height = int(pil_img.height * image_scale)
            estimated_rows = max(1, math.ceil(display_height * 0.75 / 20))

            for r in range(current_row, current_row + estimated_rows):
                ws.row_dimensions[r].height = 20

            current_row += estimated_rows + image_gap

    # -------------------------
    # Excel 导出
    # -------------------------
    def _export_to_excel(self) -> None:
        if not self.session:
            return

        image_paths = [item.image_path for item in self.session.captures]
        image_scale = self._get_image_scale()
        image_gap = self._get_image_gap()
        template_path = self.settings.get("template_path", "")
        marker = self.settings.get("screenshot_marker", "{{screenshot}}")

        if template_path and os.path.exists(template_path):
            wb = load_workbook(template_path)
            ws = wb.active
            ws.title = self._safe_sheet_name(self.sheet_name)
            self._apply_excel_zoom(ws, is_template=True)
            self._write_template_sheet_content(ws, image_paths, image_scale, marker, image_gap)
        else:
            zoom = self.settings.get("excel_zoom", None)
            zoom_scale = int(zoom) if zoom is not None else 80
            wb = Workbook()
            ws = wb.active
            ws.title = self._safe_sheet_name(self.sheet_name)
            self._write_sheet_content(
                ws=ws,
                excel_title=self.excel_title,
                image_paths=image_paths,
                image_scale=image_scale,
                zoom_scale=zoom_scale,
                image_gap=image_gap
            )

        wb.save(self.session.excel_path)

    def merge_all_excels(self) -> Optional[str]:
        """把当前 output_dir 下所有会话目录中的 Excel 整合到一个文件"""
        if not os.path.exists(self.output_dir):
            return None

        image_scale = self._get_image_scale()
        image_gap = self._get_image_gap()
        zoom = self.settings.get("excel_zoom", None)
        zoom_scale = int(zoom) if zoom is not None else 80

        merge_items = []

        for entry in os.scandir(self.output_dir):
            if not entry.is_dir():
                continue

            session_dir = entry.path

            excel_files = [
                f.path for f in os.scandir(session_dir)
                if f.is_file()
                and f.name.lower().endswith(".xlsx")
                and not f.name.startswith("~$")
            ]

            if not excel_files:
                continue

            excel_files.sort()
            source_excel = excel_files[0]

            try:
                wb = load_workbook(source_excel)
                ws = wb.active

                source_sheet_name = ws.title
                source_title = ws["A1"].value or ""

                image_paths = [
                    f.path for f in os.scandir(session_dir)
                    if f.is_file() and f.name.lower().endswith(".png")
                ]
                image_paths.sort()

                merge_items.append({
                    "sheet_name": source_sheet_name,
                    "excel_title": source_title,
                    "image_paths": image_paths,
                })
            except Exception as e:
                print(f"[整合跳过] 读取失败: {source_excel}, {e}")

        if not merge_items:
            return None

        merge_items.sort(key=lambda x: x["sheet_name"])

        merged_wb = Workbook()
        default_ws = merged_wb.active
        merged_wb.remove(default_ws)

        used_names: set = set()

        for item in merge_items:
            base_sheet_name = self._safe_sheet_name(item["sheet_name"])
            final_sheet_name = base_sheet_name
            seq = 2

            while final_sheet_name in used_names:
                suffix = f"_{seq}"
                trimmed = base_sheet_name[:31 - len(suffix)]
                final_sheet_name = f"{trimmed}{suffix}"
                seq += 1

            used_names.add(final_sheet_name)

            new_ws = merged_wb.create_sheet(title=final_sheet_name)
            self._write_sheet_content(
                ws=new_ws,
                excel_title=item["excel_title"],
                image_paths=item["image_paths"],
                image_scale=image_scale,
                zoom_scale=zoom_scale,
                image_gap=image_gap
            )

        merged_file_name = f"merged_sheets_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        merged_excel_path = os.path.join(self.output_dir, merged_file_name)
        merged_wb.save(merged_excel_path)

        return merged_excel_path


# =========================
# 设置对话框
# =========================
class SettingsDialog:
    def __init__(self, parent: tk.Tk, settings: dict, on_save) -> None:
        self.parent = parent
        self.on_save = on_save

        self.dialog = tk.Toplevel(parent)
        self.dialog.title("Settings")
        self.dialog.geometry("600x310")
        self.dialog.resizable(False, False)
        self.dialog.transient(parent)
        self.dialog.grab_set()

        self._init_vars(settings)
        self._build_ui()
        self._center_dialog()

    def _init_vars(self, settings: dict) -> None:
        self.template_var = tk.StringVar(value=settings.get("template_path", ""))
        self.marker_var = tk.StringVar(value=settings.get("screenshot_marker", "{{screenshot}}"))
        self.image_scale_var = tk.StringVar(value=str(settings.get("image_scale", 80)))
        self.image_gap_var = tk.StringVar(value=str(settings.get("image_gap", 2)))
        zoom = settings.get("excel_zoom", None)
        self.excel_zoom_auto = tk.BooleanVar(value=(zoom is None))
        self.excel_zoom_var = tk.StringVar(value=str(zoom) if zoom is not None else "80")

    def _build_ui(self) -> None:
        main = ttk.Frame(self.dialog, padding=16)
        main.pack(fill="both", expand=True)

        # --- 模板设置 ---
        tmpl_frame = ttk.LabelFrame(main, text="Excel 模板", padding=10)
        tmpl_frame.pack(fill="x", pady=(0, 10))

        ttk.Label(tmpl_frame, text="模板文件").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=4)
        ttk.Entry(tmpl_frame, textvariable=self.template_var, width=42).grid(
            row=0, column=1, sticky="ew", pady=4
        )
        ttk.Button(tmpl_frame, text="选择", command=self._choose_template).grid(
            row=0, column=2, padx=(6, 4), pady=4
        )
        ttk.Button(tmpl_frame, text="清除", command=lambda: self.template_var.set("")).grid(
            row=0, column=3, pady=4
        )

        ttk.Label(tmpl_frame, text="截图标记").grid(row=1, column=0, sticky="w", padx=(0, 8), pady=4)
        ttk.Entry(tmpl_frame, textvariable=self.marker_var, width=24).grid(
            row=1, column=1, sticky="w", pady=4
        )
        ttk.Label(
            tmpl_frame,
            text="在模板单元格中填入此文本作为截图占位符",
            foreground="gray"
        ).grid(row=1, column=2, columnspan=2, sticky="w", padx=(6, 0), pady=4)

        tmpl_frame.columnconfigure(1, weight=1)

        # --- 缩放设置 ---
        scale_frame = ttk.LabelFrame(main, text="缩放设置", padding=10)
        scale_frame.pack(fill="x", pady=(0, 10))

        ttk.Label(scale_frame, text="图片缩放").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=4)
        ttk.Entry(scale_frame, textvariable=self.image_scale_var, width=8).grid(
            row=0, column=1, sticky="w", pady=4
        )
        ttk.Label(scale_frame, text="%  （如：80 表示缩小到原图的 80%）").grid(
            row=0, column=2, sticky="w", padx=(4, 0), pady=4
        )

        ttk.Label(scale_frame, text="图片间距").grid(row=1, column=0, sticky="w", padx=(0, 8), pady=4)
        ttk.Entry(scale_frame, textvariable=self.image_gap_var, width=8).grid(
            row=1, column=1, sticky="w", pady=4
        )
        ttk.Label(scale_frame, text="行  （相邻两张图之间的空白行数）").grid(
            row=1, column=2, sticky="w", padx=(4, 0), pady=4
        )

        ttk.Label(scale_frame, text="Excel 缩放").grid(row=2, column=0, sticky="w", padx=(0, 8), pady=4)
        zoom_row = ttk.Frame(scale_frame)
        zoom_row.grid(row=2, column=1, columnspan=3, sticky="w", pady=4)
        self.excel_zoom_entry = ttk.Entry(zoom_row, textvariable=self.excel_zoom_var, width=8)
        self.excel_zoom_entry.pack(side="left")
        ttk.Label(zoom_row, text="%").pack(side="left", padx=(4, 12))
        ttk.Checkbutton(
            zoom_row,
            text="不指定（使用模板原有缩放比例）",
            variable=self.excel_zoom_auto,
            command=self._toggle_zoom
        ).pack(side="left")

        self._toggle_zoom()

        # --- 按钮 ---
        btn_frame = ttk.Frame(main)
        btn_frame.pack(fill="x", side="bottom", pady=(4, 0))
        ttk.Button(btn_frame, text="取消", command=self.dialog.destroy).pack(side="right", padx=(8, 0))
        ttk.Button(btn_frame, text="保存", command=self._save).pack(side="right")

    def _center_dialog(self) -> None:
        self.dialog.update_idletasks()
        x = self.parent.winfo_x() + (self.parent.winfo_width() - self.dialog.winfo_width()) // 2
        y = self.parent.winfo_y() + (self.parent.winfo_height() - self.dialog.winfo_height()) // 2
        self.dialog.geometry(f"+{x}+{y}")

    def _toggle_zoom(self) -> None:
        if self.excel_zoom_auto.get():
            self.excel_zoom_entry.state(["disabled"])
        else:
            self.excel_zoom_entry.state(["!disabled"])

    def _choose_template(self) -> None:
        path = filedialog.askopenfilename(
            title="选择 Excel 模板",
            filetypes=[("Excel 文件", "*.xlsx *.xlsm"), ("所有文件", "*.*")],
            parent=self.dialog
        )
        if path:
            self.template_var.set(path)

    def _save(self) -> None:
        try:
            image_scale = float(self.image_scale_var.get())
            if not (1 <= image_scale <= 200):
                raise ValueError()
        except ValueError:
            messagebox.showerror("错误", "图片缩放比例应为 1～200 之间的数字。", parent=self.dialog)
            return

        try:
            image_gap = int(self.image_gap_var.get())
        except ValueError:
            messagebox.showerror("错误", "图片间距应为整数。", parent=self.dialog)
            return

        excel_zoom = None
        if not self.excel_zoom_auto.get():
            try:
                excel_zoom = int(self.excel_zoom_var.get())
                if not (10 <= excel_zoom <= 400):
                    raise ValueError()
            except ValueError:
                messagebox.showerror("错误", "Excel 缩放比例应为 10～400 之间的整数。", parent=self.dialog)
                return

        new_settings = {
            "template_path": self.template_var.get().strip(),
            "screenshot_marker": self.marker_var.get().strip() or "{{screenshot}}",
            "image_scale": image_scale,
            "image_gap": image_gap,
            "excel_zoom": excel_zoom,
        }
        self.on_save(new_settings)
        self.dialog.destroy()


# =========================
# GUI
# =========================
class TestCaptureGUI:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("辅助测试截图工具")
        self.root.geometry("860x500")
        self.root.minsize(800, 460)

        self.app = TestCaptureApp(ui_callback=self._refresh_ui_safe)
        self.hotkey_listener = None

        self.status_var = tk.StringVar(value="未开始")
        self.capture_count_var = tk.StringVar(value="0")
        self.session_dir_var = tk.StringVar(value="")
        self.excel_path_var = tk.StringVar(value="")
        s = self.app.settings
        self.output_dir_var = tk.StringVar(value=s.get("output_dir") or self.app.output_dir)
        self.title_var = tk.StringVar(value=s.get("excel_title") or self.app.excel_title)
        self.sheet_name_var = tk.StringVar(value=s.get("sheet_name") or self.app.sheet_name)

        self.start_button: Optional[ttk.Button] = None
        self.stop_button: Optional[ttk.Button] = None
        self.reset_button: Optional[ttk.Button] = None
        self.tray_icon = None
        self.tray_thread = None
        self.is_in_tray = False
        self.toast_window = None
        self.toast_after_id = None

        self._build_ui()
        self._create_menu()
        self.refresh_ui()
        self.start_hotkeys()

        self.root.protocol("WM_DELETE_WINDOW", self.on_close)
        self.root.bind("<Unmap>", self.on_minimize)

    # -------------------------
    # 菜单
    # -------------------------
    def _create_menu(self) -> None:
        menubar = tk.Menu(self.root)

        # Settings 菜单
        settings_menu = tk.Menu(menubar, tearoff=0)
        settings_menu.add_command(label="Settings", command=self.open_settings)
        menubar.add_cascade(label="Settings", menu=settings_menu)

        # Help 菜单
        help_menu = tk.Menu(menubar, tearoff=0)
        help_menu.add_command(label="About", command=self.show_about)
        menubar.add_cascade(label="Help", menu=help_menu)

        self.root.config(menu=menubar)

    def show_about(self) -> None:
        messagebox.showinfo(
            "About",
            "Test Capture Tool v2.0\n\nAuthor: Mysoft Liu\n© 2026"
        )

    def open_settings(self) -> None:
        SettingsDialog(
            parent=self.root,
            settings=self.app.settings,
            on_save=self._on_settings_saved
        )

    def _on_settings_saved(self, new_settings: dict) -> None:
        self.app.settings = new_settings
        _save_settings(new_settings)

    # -------------------------
    # 热键
    # -------------------------
    def start_hotkeys(self) -> None:
        self.hotkey_listener = keyboard.GlobalHotKeys({
            "<ctrl>+<alt>+<f9>": lambda: self.root.after(0, self.start_recording),
            "<ctrl>+<alt>+<f10>": lambda: self.root.after(0, self.stop_recording),
        })
        self.hotkey_listener.start()

    def stop_hotkeys(self) -> None:
        if self.hotkey_listener:
            self.hotkey_listener.stop()
            self.hotkey_listener = None

    # -------------------------
    # UI 构建
    # -------------------------
    def _build_ui(self) -> None:
        main = ttk.Frame(self.root, padding=16)
        main.pack(fill="both", expand=True)

        title_label = ttk.Label(
            main,
            text="辅助测试截图工具",
            font=("Yu Gothic UI", 15, "bold")
        )
        title_label.pack(anchor="w", pady=(0, 12))

        # 输出配置
        config_frame = ttk.LabelFrame(main, text="输出配置", padding=12)
        config_frame.pack(fill="x", pady=(0, 12))

        ttk.Label(config_frame, text="Excel 标题").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=6)
        ttk.Entry(config_frame, textvariable=self.title_var, width=72).grid(row=0, column=1, columnspan=2, sticky="ew", pady=6)

        ttk.Label(config_frame, text="Sheet 名称").grid(row=1, column=0, sticky="w", padx=(0, 8), pady=6)
        ttk.Entry(config_frame, textvariable=self.sheet_name_var, width=30).grid(row=1, column=1, sticky="w", pady=6)

        ttk.Label(config_frame, text="输出目录").grid(row=2, column=0, sticky="w", padx=(0, 8), pady=6)
        ttk.Entry(config_frame, textvariable=self.output_dir_var, width=56).grid(row=2, column=1, sticky="ew", pady=6)
        ttk.Button(config_frame, text="选择目录", command=self.choose_output_dir).grid(row=2, column=2, padx=(8, 0), pady=6)

        config_frame.columnconfigure(1, weight=1)

        # 状态区域
        status_frame = ttk.LabelFrame(main, text="当前状态", padding=12)
        status_frame.pack(fill="x", pady=(0, 12))

        ttk.Label(status_frame, text="状态").grid(row=0, column=0, sticky="w", padx=(0, 12), pady=6)
        ttk.Label(status_frame, textvariable=self.status_var).grid(row=0, column=1, sticky="w", pady=6)

        ttk.Label(status_frame, text="截图数量").grid(row=1, column=0, sticky="w", padx=(0, 12), pady=6)
        ttk.Label(status_frame, textvariable=self.capture_count_var).grid(row=1, column=1, sticky="w", pady=6)

        ttk.Label(status_frame, text="会话目录").grid(row=2, column=0, sticky="nw", padx=(0, 12), pady=6)
        ttk.Label(status_frame, textvariable=self.session_dir_var, wraplength=620).grid(row=2, column=1, sticky="w", pady=6)

        ttk.Label(status_frame, text="Excel 文件").grid(row=3, column=0, sticky="nw", padx=(0, 12), pady=6)
        ttk.Label(status_frame, textvariable=self.excel_path_var, wraplength=620).grid(row=3, column=1, sticky="w", pady=6)

        # 按钮区域
        button_frame = ttk.Frame(main)
        button_frame.pack(fill="x", pady=(0, 12))

        self.start_button = ttk.Button(button_frame, text="开始录制", command=self.start_recording)
        self.start_button.pack(side="left", padx=(0, 8))

        self.stop_button = ttk.Button(button_frame, text="结束并导出 Excel", command=self.stop_recording)
        self.stop_button.pack(side="left", padx=(0, 8))

        ttk.Button(button_frame, text="打开输出目录", command=self.open_output_dir).pack(side="left", padx=(0, 8))
        ttk.Button(button_frame, text="整合文件", command=self.merge_files).pack(side="left", padx=(0, 8))

        self.reset_button = ttk.Button(button_frame, text="清空状态", command=self.reset_session)
        self.reset_button.pack(side="left")

        # 说明区域
        help_frame = ttk.LabelFrame(main, text="使用说明", padding=12)
        help_frame.pack(fill="both", expand=True)

        help_text = (
            "开始方式：\n"
            '  1. 点击"开始录制"按钮\n'
            "  2. 或按快捷键 Ctrl + Alt + F9\n\n"
            "结束方式：\n"
            '  1. 点击"结束并导出 Excel"按钮\n'
            "  2. 或按快捷键 Ctrl + Alt + F10\n\n"
            "截图方式：\n"
            "  使用 Alt + PrintScreen 截取当前活动窗口。\n\n"
            "说明：\n"
            "  录制开始后，每次用 Alt + PrintScreen 截图，程序会自动保存图片。\n"
            "  结束后会自动生成 Excel，并直接打开保存好的 Excel 文件。"
        )
        ttk.Label(help_frame, text=help_text, justify="left").pack(anchor="w")

    # -------------------------
    # 设置
    # -------------------------
    def apply_settings(self) -> None:
        self.app.excel_title = self.title_var.get().strip() or "截图结果"
        self.app.sheet_name = self.sheet_name_var.get().strip() or "0001"
        self.app.output_dir = self.output_dir_var.get().strip() or self.app.base_output_dir
        os.makedirs(self.app.output_dir, exist_ok=True)
        self.app.settings["excel_title"] = self.app.excel_title
        self.app.settings["sheet_name"] = self.app.sheet_name
        self.app.settings["output_dir"] = self.app.output_dir
        _save_settings(self.app.settings)

    def choose_output_dir(self) -> None:
        selected = filedialog.askdirectory(initialdir=self.output_dir_var.get() or os.getcwd())
        if selected:
            self.output_dir_var.set(selected)

    # -------------------------
    # 操作
    # -------------------------
    def start_recording(self) -> None:
        try:
            self.apply_settings()
            started = self.app.start_recording()
            self.refresh_ui()
            if not started:
                messagebox.showinfo("提示", "当前已经在录制中。")
        except Exception as e:
            messagebox.showerror("错误", f"开始录制失败：\n{e}")

    def stop_recording(self) -> None:
        try:
            stopped = self.app.stop_recording()
            self.refresh_ui()

            if not stopped:
                messagebox.showinfo("提示", "当前没有录制中的会话。")
                return

            if self.app.session and os.path.exists(self.app.session.excel_path):
                try:
                    os.startfile(self.app.session.excel_path)
                except Exception as open_err:
                    messagebox.showwarning(
                        "提示",
                        f"Excel 已生成，但自动打开失败：\n{open_err}\n\n文件位置：\n{self.app.session.excel_path}"
                    )
        except Exception as e:
            messagebox.showerror("错误", f"结束录制失败：\n{e}")

    def open_output_dir(self) -> None:
        path = ""
        if self.app.session and self.app.session.session_dir:
            path = self.app.session.session_dir
        else:
            path = self.output_dir_var.get().strip() or self.app.base_output_dir

        if not os.path.exists(path):
            os.makedirs(path, exist_ok=True)

        try:
            os.startfile(path)
        except Exception as e:
            messagebox.showerror("错误", f"打开目录失败：\n{e}")

    def merge_files(self) -> None:
        try:
            self.apply_settings()
            merged_excel_path = self.app.merge_all_excels()

            if not merged_excel_path:
                messagebox.showinfo("提示", "当前输出目录下没有可整合的 Excel 文件。")
                return

            try:
                os.startfile(merged_excel_path)
            except Exception as open_err:
                messagebox.showwarning(
                    "提示",
                    f"整合文件已生成，但自动打开失败：\n{open_err}\n\n文件位置：\n{merged_excel_path}"
                )
        except Exception as e:
            messagebox.showerror("错误", f"整合文件失败：\n{e}")

    def reset_session(self) -> None:
        if self.app.session and self.app.session.is_recording:
            messagebox.showwarning("提示", "录制中不能清空，请先结束录制。")
            return

        self.app.reset_session()
        self.refresh_ui()

    # -------------------------
    # 刷新UI
    # -------------------------
    def _refresh_ui_safe(self, event_name=None, payload=None) -> None:
        if event_name == "capture_saved":
            self.root.after(0, lambda: self.on_capture_saved(payload))
            return

        self.root.after(0, self.refresh_ui)

    def on_capture_saved(self, item: Optional[CaptureItem]) -> None:
        self.refresh_ui()

        if item is None:
            return

        message = f"第 {item.index} 张截图已保存"

        if self.is_in_tray and self.tray_icon:
            try:
                self.tray_icon.notify(message, "Test Capture Tool")
                return
            except Exception:
                pass

        self.show_toast(message)

    def show_toast(self, message: str, duration_ms: int = 1800) -> None:
        try:
            if self.toast_after_id is not None:
                self.root.after_cancel(self.toast_after_id)
                self.toast_after_id = None
        except Exception:
            self.toast_after_id = None

        try:
            if self.toast_window is not None and self.toast_window.winfo_exists():
                self.toast_window.destroy()
        except Exception:
            pass

        toast = tk.Toplevel(self.root)
        self.toast_window = toast
        toast.overrideredirect(True)
        toast.attributes("-topmost", True)
        toast.configure(bg="#2b2b2b")

        frame = tk.Frame(toast, bg="#2b2b2b", bd=1, relief="solid")
        frame.pack(fill="both", expand=True)

        title_label = tk.Label(
            frame,
            text="Test Capture Tool",
            font=("Yu Gothic UI", 10, "bold"),
            fg="white",
            bg="#2b2b2b",
            anchor="w"
        )
        title_label.pack(fill="x", padx=12, pady=(10, 2))

        message_label = tk.Label(
            frame,
            text=message,
            font=("Yu Gothic UI", 10),
            fg="white",
            bg="#2b2b2b",
            justify="left",
            anchor="w"
        )
        message_label.pack(fill="x", padx=12, pady=(0, 10))

        toast.update_idletasks()

        toast_width = max(260, toast.winfo_reqwidth())
        toast_height = max(80, toast.winfo_reqheight())

        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()

        x = screen_width - toast_width - 20
        y = screen_height - toast_height - 60
        toast.geometry(f"{toast_width}x{toast_height}+{x}+{y}")

        def close_toast() -> None:
            try:
                if self.toast_window is not None and self.toast_window.winfo_exists():
                    self.toast_window.destroy()
            except Exception:
                pass
            finally:
                self.toast_window = None
                self.toast_after_id = None

        self.toast_after_id = self.root.after(duration_ms, close_toast)

    def create_tray_image(self) -> Image.Image:
        size = 64
        image = Image.new("RGBA", (size, size), (52, 120, 246, 255))
        draw = ImageDraw.Draw(image)
        draw.rectangle((16, 14, 48, 50), fill=(255, 255, 255, 255))
        draw.rectangle((21, 10, 43, 18), fill=(255, 255, 255, 255))
        draw.rectangle((22, 23, 42, 27), fill=(52, 120, 246, 255))
        draw.rectangle((22, 31, 42, 35), fill=(52, 120, 246, 255))
        draw.rectangle((22, 39, 34, 43), fill=(52, 120, 246, 255))
        return image

    def show_window_from_tray(self, icon=None, item=None) -> None:
        self.root.after(0, self.restore_from_tray)

    def restore_from_tray(self) -> None:
        if self.tray_icon:
            try:
                self.tray_icon.stop()
            except Exception:
                pass
            self.tray_icon = None

        self.is_in_tray = False
        try:
            self.root.iconify()
            self.root.deiconify()
        except Exception:
            pass
        self.root.deiconify()
        self.root.after(0, self.root.lift)
        self.root.after(0, self.root.focus_force)

    def quit_from_tray(self, icon=None, item=None) -> None:
        self.root.after(0, self.on_close)

    def _run_tray_icon(self) -> None:
        if pystray is None:
            return

        menu = pystray.Menu(
            pystray.MenuItem("打开", self.show_window_from_tray, default=True),
            pystray.MenuItem("退出", self.quit_from_tray)
        )

        self.tray_icon = pystray.Icon(
            "test_capture_to_excel",
            self.create_tray_image(),
            "Test Capture Tool",
            menu
        )
        self.tray_icon.run()

    def minimize_to_tray(self) -> None:
        if pystray is None or self.is_in_tray:
            return

        self.is_in_tray = True
        self.root.withdraw()

        self.tray_thread = threading.Thread(target=self._run_tray_icon, daemon=True)
        self.tray_thread.start()

    def on_minimize(self, event=None) -> None:
        try:
            if self.root.state() == "iconic":
                self.minimize_to_tray()
        except Exception:
            pass

    def refresh_ui(self) -> None:
        session = self.app.session

        if session and session.is_recording:
            self.status_var.set("录制中")
        elif session:
            self.status_var.set("已结束")
        else:
            self.status_var.set("未开始")

        if session:
            self.capture_count_var.set(str(len(session.captures)))
            self.session_dir_var.set(session.session_dir)
            self.excel_path_var.set(session.excel_path)
        else:
            self.capture_count_var.set("0")
            self.session_dir_var.set("")
            self.excel_path_var.set("")

        if self.start_button and self.stop_button and self.reset_button:
            if session and session.is_recording:
                self.start_button.state(["disabled"])
                self.stop_button.state(["!disabled"])
                self.reset_button.state(["disabled"])
            else:
                self.start_button.state(["!disabled"])
                self.stop_button.state(["!disabled"])
                self.reset_button.state(["!disabled"])

    # -------------------------
    # 关闭
    # -------------------------
    def on_close(self) -> None:
        try:
            self.app.stop_event.set()
            self.stop_hotkeys()
            if self.tray_icon:
                try:
                    self.tray_icon.stop()
                except Exception:
                    pass
                self.tray_icon = None
        finally:
            try:
                if self.toast_after_id is not None:
                    self.root.after_cancel(self.toast_after_id)
                    self.toast_after_id = None
            except Exception:
                self.toast_after_id = None

            try:
                if self.toast_window is not None and self.toast_window.winfo_exists():
                    self.toast_window.destroy()
            except Exception:
                pass
            finally:
                self.toast_window = None

            self.root.destroy()
            sys.exit(0)


if __name__ == "__main__":
    root = tk.Tk()
    style = ttk.Style()
    try:
        style.theme_use("vista")
    except tk.TclError:
        pass

    gui = TestCaptureGUI(root)
    root.mainloop()
